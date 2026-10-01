import re, sys
sys.path.insert(0, '/tmp/rt2')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from locate import loc

MS = docx.Document('Manuscript_Revised_Clean.docx')
MS_PARAS = []
for p in MS.element.body.iter(qn('w:p')):
    MS_PARAS.append(''.join(t.text or '' for t in p.iter(qn('w:t'))))
TP = docx.Document('Title_Page_Revised_Clean.docx')
TP_PARAS = [p.text for p in TP.paragraphs]
S1 = docx.Document('S1_Supplementary_Interview_Guide_Revised_Clean.docx')
S1_PARAS = [''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in S1.element.body.iter(qn('w:p'))]

USED = []
def Q(start, end=None, src='ms'):
    """Exact substring of the revised file from start..end (inclusive)."""
    paras = {'ms': MS_PARAS, 'tp': TP_PARAS, 's1': S1_PARAS}[src]
    for t in paras:
        i = t.find(start)
        if i >= 0:
            if end is None:
                return t[i:].rstrip()
            j = t.find(end, i)
            assert j >= 0, ('end not found', end)
            s = t[i:j + len(end)]
            USED.append(s)
            return s
    raise AssertionError('quote start not found: ' + start)

def L(start, end=None):
    r = loc(start, end)
    assert r, ('locate failed', start)
    return r

doc = docx.Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.54); s.left_margin = s.right_margin = Cm(2.54)


PPR_ORDER = ['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd','snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc','textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange']
def ppr_put(p, el):
    pPr = p._p.get_or_add_pPr()
    name = el.tag.split('}')[1]
    existing = pPr.find(qn('w:' + name))
    if existing is not None:
        if name == 'pBdr':
            for c in el: existing.append(c)
            return
        pPr.remove(existing)
    idx = PPR_ORDER.index(name)
    for i, c in enumerate(list(pPr)):
        cn = c.tag.split('}')[1]
        if cn in PPR_ORDER and PPR_ORDER.index(cn) > idx:
            c.addprevious(el); return
    pPr.append(el)

def shade(p, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill); ppr_put(p, shd)

def border_left(p, color='1F3864'):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); l = OxmlElement('w:left')
    for k, v in (('val', 'single'), ('sz', '18'), ('space', '8'), ('color', color)): l.set(qn('w:' + k), v)
    b.append(l); ppr_put(p, b)

PH = re.compile(r'(\[AUTHOR:[^\]]*\])')
def runs(p, text, bold=False, italic=False, size=None, color=None):
    for part in PH.split(text):
        if not part: continue
        r = p.add_run(part); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if PH.fullmatch(part):
            r.font.highlight_color = 7  # yellow
    return p

def para(text='', bold=False, italic=False, align=None, size=None, after=None, color=None):
    p = doc.add_paragraph(); runs(p, text, bold, italic, size, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def heading(text):
    p = para(text, bold=True, size=13, color='1F3864', after=6)
    p.paragraph_format.space_before = Pt(14)
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '6'), ('space', '1'), ('color', '1F3864')): bt.set(qn('w:' + k), v)
    b.append(bt); ppr_put(p, b)

def comment(label, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    runs(p, label + ' ', bold=True, color='1F3864'); runs(p, text, italic=True)
    shade(p, 'EEF3FA')

def response(text):
    p = doc.add_paragraph(); runs(p, 'Response: ', bold=True); runs(p, text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def revised(where, text):
    p = doc.add_paragraph(); runs(p, f'Revised text ({where}):', bold=True, italic=True, size=11)
    p.paragraph_format.space_after = Pt(2)
    q = doc.add_paragraph(); runs(q, '“' + text + '”', size=11)
    q.paragraph_format.left_indent = Cm(0.8); border_left(q); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); runs(p, text); p.paragraph_format.space_after = Pt(3)

# ---------------------------------------------------------------- header
box = para('BEFORE SUBMISSION (delete this box): complete every yellow [AUTHOR: …] placeholder here and in the manuscript, title page, and Supplementary Files S2–S3, and insert the redrawn Figure 2. See Final_Audit_Report.md.', bold=True, size=10)
shade(box, 'FFF2CC')
para('Response to the Editor and Reviewers', bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, color='1F3864')
for k, v in (('Manuscript ID', '7545151'),
             ('Title', 'Between Alert and Action: Nurses’ Experiences of AI-Assisted Early Warning, Clinical Judgment, and Patient Safety in Critical Care'),
             ('Journal', 'Nursing & Health Sciences')):
    p = doc.add_paragraph(); runs(p, k + ': ', bold=True); runs(p, v); p.paragraph_format.space_after = Pt(0)
para('')
para('Dear Editor,')
para('We thank the editor and both reviewers for their constructive comments. Each comment is answered below with the revised text and its location (page and line numbers refer to the highlighted manuscript). All changes are highlighted in yellow; a clean copy is also supplied.', align=WD_ALIGN_PARAGRAPH.JUSTIFY)
para('In addition, we checked the manuscript for internal consistency and made these corrections, none of which alters the findings:', align=WD_ALIGN_PARAGRAPH.JUSTIFY)
for b_ in [
    'Participant counts now agree with the participant table: female n = 12 (52%; previously 13, 57%); five or six participants per site (previously “five to seven”); nurses with five or fewer years of experience n = 6 (previously n = 7). A sentence repeating the experience range in Section 3.1 was deleted.',
    'References: two duplicates merged (former 35 and 36 = 12 and 26); misattributed citations corrected (e.g., Braun and Clarke, Thorne); one unrelated citation removed; six entries reformatted; 11 references added; list renumbered (77 references).',
    'New Table 1 added; former Tables 1 and 2 are now Tables 2 and 3. Reversed age ranges in the participant table (a right-to-left formatting error) corrected.',
    'Overstated wording moderated in the Abstract conclusion, Key points, and Discussion.',
]:
    bullet(b_)
para('Yours sincerely,', after=0)
para('Osama Mohamed Elsayed Ramadan, on behalf of all authors')

# ---------------------------------------------------------------- Editor
heading('EDITOR')
comment('Comment E1.', 'Characterization of the systems and participants’ exposure. If possible, describe the systems used at each hospital, including their purpose, the component that qualifies them as AI-assisted, the alerts generated, and their integration into nursing workflows. Distinguish these systems from conventional monitoring alarms or rule-based decision-support systems. Please also report how long the systems had been implemented and participants’ experience using them, drawing on the information collected through the demographic form. Explain how differences in technology, exposure, training, and response protocols across hospitals informed the interpretation of the findings. A concise table may help present this information while preserving institutional anonymity. If this distinction cannot be made, please explain why and discuss the implications for interpreting the findings in the Discussion section.')
response('(a) Definition and distinction. We added a working definition that separates AI-assisted systems from monitor alarms and rule-based scores. The same distinction now opens the Introduction, and the setting, inclusion criterion, and participant description use the defined term consistently (Sections 2.3, 2.4.1, and 3.1).')
revised(f'Section 2.3, {L("For this study, AI-assisted early warning", "preserve institutional anonymity.")}', Q('For this study, AI-assisted early warning', 'preserve institutional anonymity.'))
response('(b, c) Systems and exposure. A new Table 1 reports, for each site, the system’s purpose and AI component, alert types, workflow integration, implementation period, training, and response protocol. It also reports participants’ duration of use, training, and prior experience, taken from Section 3 of the demographic form, which is now described in Section 2.5. Vendor names are withheld to preserve anonymity. [AUTHOR: complete Table 1, then delete this note.]')
revised(f'Section 2.5, {L("Before each interview, participants completed", "(Supplementary File S1).")}', Q('Before each interview, participants completed', '(Supplementary File S1).'))
response('(d) Use in analysis. Site differences were treated as interpretive context during analysis.')
revised(f'Section 2.8, {L("Site-level differences (Table 1)", "described at Hospital A.")}', Q('Site-level differences (Table 1)', 'described at Hospital A.'))
response('(e) Implications. The Limitations now separate cross-site patterns from site-specific ones. They also state that the findings concern locally implemented systems, and that alert burden may partly reflect conventional alarms.')
revised(f'Section 4.6, {L("The four sites used different systems", "system audit data.")}', Q('The four sites used different systems', 'system audit data.'))

comment('Comment E2.', 'Transparency in theme development. The supplementary interview guide labels its sections “Themes 1–4,” closely corresponding to the final themes. Please clarify whether these labels were present during data collection or added retrospectively for presentation. Explain the relationship between the orienting conceptual framework, interview questions, initial codes, and final themes, distinguishing predefined areas of inquiry from interpretations developed through analysis. Please also clarify whether the analysis was deductive, inductive, or a combination of both.')
response('(a) The theme labels were added retrospectively, when the supplementary file was prepared after analysis. During data collection, sections were labelled by area of inquiry. We have restored those labels in S1, added an explanatory note there, and stated this in Section 2.6. [AUTHOR: confirm; see audit item A1.]')
revised(f'Section 2.6, {L("These topics were organised into four areas", "(Section 2.8).")}', Q('These topics were organised into four areas', '(Section 2.8).'))
response('(b, c) The analysis combined both approaches: the framework set the areas of inquiry, while codes and themes were developed inductively. Section 2.8 now explains this and names the interpretations that were not predefined. New Supplementary File S2 traces framework concept → area of inquiry → illustrative codes → theme for each theme. The correspondence between the guide and the theme structure is acknowledged as a limitation (Section 4.6).')
revised(f'Section 2.8, {L("The analysis combined deductive and inductive", "and final themes.")}', Q('The analysis combined deductive and inductive', 'and final themes.'))

comment('Comment E3.', 'Disclosure of AI use. Please specify whether AI tools were used in preparing the manuscript. If they were used for drafting, language editing, grammar correction, or similar purposes, this information must be included in the Acknowledgements section on the title page. If AI tools were used to create or enhance Figures 1 and 2, this information must also be disclosed in the respective figure legends.')
response('The following statement has been added to the Acknowledgements on the title page. Both figure legends now state that no AI tools were used to create or enhance the figures. [AUTHOR: if an AI assistant was used to help prepare this revision, disclose it here and on the title page; see audit item B0.]')
revised('Title page, Acknowledgements', Q('The authors used Grammarly and QuillBot', 'full responsibility for its content.', src='tp'))

comment('Comment E4.', 'These revisions are essential to establish what technology participants experienced and how the analysis generated its interpretive contribution. Please submit a revised manuscript with a point-by-point response explaining how each concern has been addressed.')
response('Addressed under E1 (technology: Table 1, Section 2.3) and E2 (analysis: Section 2.8, Supplementary File S2).')

# ---------------------------------------------------------------- Reviewer 1
heading('REVIEWER 1')
comment('Comment 1.1.', 'The manuscript submitted by the authors is very well presented and comprehensively addresses the core aspects of research. The study is ethically sound, scientifically rigorous, and methodologically robust, with clear and technically accurate reporting of the methods and results. The scope of the manuscript is highly relevant to clinical practice, education, and future decision-making. The findings are presented in a balanced manner, offering a comprehensive understanding of participants’ perspectives. Importantly, the authors have also appropriately acknowledged and discussed the limitations of the study, reflecting transparency and critical reflection in reporting the research. Furthermore, the detailed presentation of the study has the potential to encourage similar research in other contexts and settings. Overall, the manuscript is reader-friendly, engaging, and contributes meaningfully to the existing body of knowledge.')
response('We thank the reviewer for this positive assessment. No changes were requested.')

# ---------------------------------------------------------------- Reviewer 2
heading('REVIEWER 2')
comment('General comment.', 'Thank you to the authors for this well-presented manuscript which presents research that highlights the complexity of integrating clinical judgement with technology use in critical care environments. The study is rigorous and adheres well to the chosen methodology. … The findings are informative and elucidate the necessity for critical thinking and following intuitive and experiential judgments alongside adapting to the technology. Reporting on contextual differences is an added strength of the multi-site study. I have provided some feedback for consideration, mainly relating to clarifying phrasing and supporting references. I also have a concern about table 1 and potential identification of participants. Please see below for more details.')
response('Thank you. Each point is answered below.')
comment('Comment 2.1.', 'Please consider adding a COREQ checklist to the submission to support reporting quality.')
response('The completed checklist is provided as Supplementary File S3 and cited in Section 2.1 and the Abstract. Three items not previously reported (non-participants present, repeat interviews, and transcript return) were added to Sections 2.5 and 2.10.')
revised(f'Section 2.1, {L("Reporting followed the Consolidated Criteria", "Supplementary File S3.")}', Q('Reporting followed the Consolidated Criteria', 'Supplementary File S3.'))

comment('Comment 2.2.', 'Line 82, page 5 - An additional sentence to outline what Vision 2030 is would provide more context to this statement.')
response('Added.')
revised(f'Introduction, {L("Vision 2030, launched in 2016", "digital services [31].")}', Q('Vision 2030, launched in 2016', 'digital services [31].'))

comment('Comment 2.3.', 'Is there any underpinning reference for figure 1? If AI was used to develop the figure, please disclose this or add in any supporting references used.')
response('The framework is the authors’ own. Its sources are now cited: Tanner’s clinical judgment model [43], Sittig and Singh’s sociotechnical model [44], and the literature on trust in automation and alarm fatigue [11, 45, 46]. The legend states that no AI tools were used.')
revised(f'Section 2.2, {L("The framework was developed by the authors from", "contextual influences.")}', Q('The framework was developed by the authors from', 'contextual influences.'))

comment('Comment 2.4.', 'Did the researcher/ interviewer have any relationship with any of the settings? For example, have they previously worked in ICU? Please consider adding this information to the reflexivity statements on page 11; section 2.10.')
response('Added to Section 2.10.')
revised(f'Section 2.10, {L("The principal investigator is a male", "reviewed in team meetings.")}', Q('The principal investigator is a male', 'reviewed in team meetings.'))

comment('Comment 2.5.', 'I have some concerns that the confidentiality of the participants may be at risk though the information in table 1. For example, if there is only one male, Filipino ICU nurse working in hospital site A they would be able to be identified. I think the nationalities could be removed from this table, but the aggregated text could stay in (line 288, page 12).')
response('Agreed. Nationality has been removed from the participant table (now Table 2) and is reported only in aggregate in the text. To further reduce identifiability, experience is banded in the table and in quotation attributions, and the single doctoral qualification is grouped as “Postgraduate”.')
revised(f'Table 2 note, {L("Note: To protect participant confidentiality", "doctoral degree (n = 1).")}', Q('Note: To protect participant confidentiality', 'doctoral degree (n = 1).'))

comment('Comment 2.6.', 'Figure 2 is very difficult to read. I appreciate the interconnectedness of all the concepts but please revise the figure to ensure all text boxes are easier to read.')
response('Figure 2 has been redrawn: the background network was removed, text is at least 9 pt, each theme is one block containing its subthemes, labels match Table 3, and a typographical error was corrected. [AUTHOR: redraw per audit item B2, confirm that each statement holds, then delete this note.]')

comment('Comment 2.7.', 'Lines 569-571 page 22; please provide a reference to the earlier work mentioned here. Also, the statements on lines 572 -575 would benefit from additional referencing support.')
response('The earlier work (technology-acceptance research) is now cited [9, 36, 62–64]. The statements that followed are supported with references on the Saudi nursing workforce and expatriate nurses’ adjustment [65, 66] and on experiential learning [67].')
revised(f'Section 4.1, {L("Earlier work has largely treated technology acceptance", "structured onboarding.")}', Q('Earlier work has largely treated technology acceptance', 'structured onboarding.'))

comment('Comment 2.8.', 'Line 604 -please provide a reference for ‘sociotechnical perspectives.’')
response('References added [44, 72].')
revised(f'Section 4.3, {L("This interpretation aligns with sociotechnical", "[44, 72].")}', Q('This interpretation aligns with sociotechnical', '[44, 72].'))

comment('Comment 2.9.', 'Line 608 – For clarity, rather than starting the sentence with ‘This suggests…’, consider rephrasing to, “These findings suggest…’')
response('Changed as suggested. Other sentences opening with “This suggests” or “This highlights” were also revised (Findings subtheme 3.1; Discussion Sections 4.1, 4.2, and 4.4).')
revised(f'Section 4.3, {L("These findings suggest that the safety", "most vulnerable.")}', Q('These findings suggest that the safety', 'most vulnerable.'))

comment('Comment 2.10.', 'Line 621 – again, please rephrase ‘This suggests..’ as it can be ambiguous.')
response('Rephrased.')
revised(f'Section 4.4, {L("Taken together, these accounts suggest", "particular context [45].")}', Q('Taken together, these accounts suggest', 'particular context [45].'))

comment('Comment 2.11.', 'It is interesting that the age range of participants overall was 26 – 51 years. This may be a factor to add into the discussion and/or limitations. Does this reflect the general demographic of ICU nurses overall in the specific contexts?')
response('Added to the Limitations. Unit-level workforce data were not available, so representativeness cannot be confirmed; the text says so.')
revised(f'Section 4.6, {L("Participants were aged 26 to 51 years.", "may differ.")}', Q('Participants were aged 26 to 51 years.', 'may differ.'))

alltxt = '\n'.join(MS_PARAS)
for c in ['[43]', '[44]', '[11, 45, 46]', '[9, 36, 62, 63]', '[64]', '[65]', '[66]', '[67]', '[44, 72]', '[31]', '[45]']:
    assert c in alltxt, ('citation in letter not in manuscript', c)
z = doc.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'), '100')
doc.save('Response_to_Editor_and_Reviewers.docx')
print('letter saved; quotes used:', len(USED))
